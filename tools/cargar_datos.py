#!/usr/bin/env python3
"""Carga datos de referencia y de listas en la base security_db (MySQL).

Comandos (se pueden combinar):
  listas       Catalogo de listas y sus 6 grupos de colorimetria (Excel "Listas y Colorimetria.xlsx").
               Corrige tambien la marca es_pep de las listas de Personas Expuestas Politicamente.
  sancionadas  Sujetos obligados sancionados por la SBS (Excel "Lista de empresas sancionadas.xlsx").
  politicos    Figuras politicas y sus casos en prensa (JSON exportado del ML o API del ML).
  todo         listas + sancionadas (+ politicos si se indica --politicos-json o --ml-url).

SEGURO POR DEFECTO: sin --aplicar solo SIMULA (hace todo dentro de una transaccion y la deshace) y muestra
que cambiaria. Es idempotente: se puede ejecutar varias veces; busca por claves naturales (nombre de la lista,
tipo + numero de documento) y nunca borra nada. Sobre la base compartida (Aiven) haz antes un respaldo:
ver tools/README.md.

Los datos personales (nombres, documentos) NO viven en este repositorio: se leen de tus archivos en el momento.
"""

import argparse
import json
import os
import re
import sys
import unicodedata
import urllib.parse
import urllib.request
from datetime import date, datetime

import openpyxl
import pymysql

GRUPOS = [  # (orden, nombre, color) - igual que las bases del equipo
    (1, "Listas Restrictivas", "Vinotinto"),
    (2, "Listas Asociadas a LA/FT o Corrupcion (Penal)", "Rojo"),
    (3, "Listas Asociadas a LA/FT o Corrupcion (Administrativo)", "Naranja"),
    (4, "Sanciones Administrativas", "Amarillo"),
    (5, "Listas de Afectacion Financiera", "Amarillo"),
    (6, "Listas Informativas y PEPs", "Verde"),
]
LISTA_SBS = "Peru - Personas sancionadas por la Superintendencia de Banca y Seguros"
LISTA_NOTICIAS = "Noticias"
PARTICULAS = {"DE", "DEL", "LA", "LAS", "LOS", "SAN", "SANTA", "VDA", "DA", "DI", "VAN", "VON", "Y"}


# ----------------------------------------------------------------------------- utilidades

def sin_tildes(t: str) -> str:
    return unicodedata.normalize("NFD", str(t)).encode("ascii", "ignore").decode()


def clave(t: str) -> str:
    return re.sub(r"\s+", " ", sin_tildes(t)).strip().upper()


def leer_env(ruta: str) -> dict:
    datos = {}
    if os.path.exists(ruta):
        for linea in open(ruta, encoding="utf-8"):
            m = re.match(r"\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*?)\s*$", linea)
            if m and not linea.lstrip().startswith("#"):
                datos[m.group(1)] = m.group(2)
    return datos


def conectar(args):
    env = leer_env(args.env)
    url = args.db_url or os.environ.get("DB_URL") or env.get("DB_URL") or "jdbc:mysql://localhost:3307/security_db"
    usuario = args.usuario or os.environ.get("DB_USER") or env.get("DB_USER") or "root"
    clave_db = args.clave if args.clave is not None else os.environ.get("DB_PASSWORD", env.get("DB_PASSWORD", "root"))
    m = re.match(r"jdbc:mysql://([^:/]+)(?::(\d+))?/([^?]+)(?:\?(.*))?", url)
    if not m:
        sys.exit(f"DB_URL no reconocida: {url}")
    host, puerto, base, opciones = m.group(1), int(m.group(2) or 3306), m.group(3), m.group(4) or ""
    ssl = {"ssl": True} if re.search(r"ssl-?mode=(REQUIRED|VERIFY)", opciones, re.I) else None
    print(f"Conectando a {host}:{puerto}/{base} como {usuario} (SSL: {'si' if ssl else 'no'})")
    return pymysql.connect(host=host, port=puerto, user=usuario, password=clave_db, database=base,
                           charset="utf8mb4", ssl=ssl, autocommit=False)


class Base:
    """Pequeno envoltorio con contadores para el informe final."""

    def __init__(self, conexion):
        self.c = conexion
        self.cur = conexion.cursor()
        self.stats = {}

    def uno(self, sql, params=()):
        self.cur.execute(sql, params)
        fila = self.cur.fetchone()
        return fila[0] if fila else None

    def ejecutar(self, sql, params=()):
        self.cur.execute(sql, params)
        return self.cur.lastrowid

    def contar(self, k, n=1):
        self.stats[k] = self.stats.get(k, 0) + n


# ----------------------------------------------------------------------------- catalogos base

def id_pais(db: Base, nombre="Peru"):
    return db.uno("SELECT id FROM pais WHERE nombre = %s", (nombre,))


def id_tipo_doc(db: Base, nombre, descripcion=None):
    i = db.uno("SELECT id FROM tipo_documento WHERE nombre = %s", (nombre,))
    if i is None:
        i = db.ejecutar("INSERT INTO tipo_documento (nombre, descripcion) VALUES (%s, %s)", (nombre, descripcion))
        db.contar("tipo_documento nuevos")
    return i


def id_grupo(db: Base, nombre):
    for gid, gnombre in _grupos(db):
        if clave(gnombre) == clave(nombre):
            return gid
    return None


def _grupos(db: Base):
    db.cur.execute("SELECT id, nombre FROM lista_grupo")
    return db.cur.fetchall()


def id_lista(db: Base, nombre):
    db.cur.execute("SELECT id, nombre FROM tipo_lista")
    for lid, lnombre in db.cur.fetchall():
        if clave(lnombre or "") == clave(nombre):
            return lid
    return None


def id_lista_like(db: Base, patron_regex):
    db.cur.execute("SELECT id, nombre FROM tipo_lista ORDER BY id")
    for lid, lnombre in db.cur.fetchall():
        if re.search(patron_regex, sin_tildes(lnombre or ""), re.I):
            return lid
    return None


# ----------------------------------------------------------------------------- listas

def cargar_listas(db: Base, carpeta):
    ruta = _archivo(carpeta, "Listas y Colorimetría.xlsx")
    print(f"\n== Listas: {os.path.basename(ruta)}")
    for orden, nombre, color in GRUPOS:
        if id_grupo(db, nombre) is None:
            db.ejecutar("INSERT INTO lista_grupo (orden, nombre, color) VALUES (%s, %s, %s)", (orden, nombre, color))
            db.contar("grupos nuevos")
    grupos = {orden: id_grupo(db, nombre) for orden, nombre, _ in GRUPOS}

    ws = openpyxl.load_workbook(ruta, read_only=True, data_only=True)["Table 1"]
    filas = [r for r in ws.iter_rows(values_only=True, min_row=26, max_col=9) if isinstance(r[1], (int, float)) and r[2]]
    for r in filas:
        nombre = re.sub(r"\s+", " ", str(r[2])).strip()
        descripcion, fuente, alcance, grupo, periodicidad = r[4], r[5], r[6], r[7], r[8]
        es_pep = 1 if re.search(r"Expuest|Peps", sin_tildes(nombre), re.I) else 0
        existente = id_lista(db, nombre)
        if existente is None:
            db.ejecutar(
                "INSERT INTO tipo_lista (nombre, descripcion, fuente, alcance, periodicidad, es_pep, id_grupo) "
                "VALUES (%s, %s, %s, %s, %s, %s, %s)",
                (nombre, descripcion, fuente, alcance, periodicidad, es_pep, grupos.get(int(grupo)) if grupo else None))
            db.contar("listas nuevas")
        else:
            cambios = db.ejecutar("UPDATE tipo_lista SET es_pep = %s WHERE id = %s AND (es_pep IS NULL OR es_pep <> %s)",
                                  (es_pep, existente, es_pep))
            if db.cur.rowcount:
                db.contar("listas con es_pep corregido")
    # Las listas demo "Personas Expuestas Politicamente" (si existen) tambien son PEP.
    db.ejecutar("UPDATE tipo_lista SET es_pep = 1 WHERE nombre LIKE %s AND (es_pep IS NULL OR es_pep = 0)", ("%Expuestas Pol%",))
    if db.cur.rowcount:
        db.contar("listas demo PEP corregidas", db.cur.rowcount)
    print(f"   catalogo del Excel: {len(filas)} listas")


# ----------------------------------------------------------------------------- sancionadas

def dividir_nombre(completo: str):
    """'QUEVEDO MALMACEDA WILVEDER' -> (nombres, ape_pat, ape_mat). Respeta particulas (DE LA CRUZ)."""
    t = completo.split()

    def apellido(i):
        j = i
        while j < len(t) - 1 and t[j].upper() in PARTICULAS:
            j += 1
        return " ".join(t[i:j + 1]), j + 1

    if len(t) < 3:
        return " ".join(t[1:]) or None, t[0] if t else None, None
    ap_pat, i = apellido(0)
    ap_mat, i = apellido(i) if i < len(t) - 1 else (None, i)
    nombres = " ".join(t[i:]) or None
    return nombres, ap_pat, ap_mat


def cargar_sancionadas(db: Base, carpeta):
    ruta = _archivo(carpeta, "Lista de empresas sancionadas.xlsx")
    print(f"\n== Sancionadas SBS: {os.path.basename(ruta)}")
    lista = id_lista(db, LISTA_SBS)
    if lista is None:
        sys.exit(f"Falta la lista '{LISTA_SBS}': ejecuta primero el comando 'listas'.")
    ruc, pais = id_tipo_doc(db, "RUC", "Registro Unico de Contribuyentes"), id_pais(db)
    ws = openpyxl.load_workbook(ruta, read_only=True, data_only=True).worksheets[0]
    filas = [r for r in ws.iter_rows(values_only=True, min_row=2, max_col=7) if r[2]]
    hoy = date.today()
    for _, nombre, doc, rubro, sancion, anio, contencioso in filas:
        doc, nombre = str(doc).strip(), re.sub(r"\s+", " ", str(nombre)).strip()
        natural = doc.startswith("10")
        eid = db.uno("SELECT id FROM entidades WHERE id_tipo_documento = %s AND documento = %s", (ruc, doc))
        if eid is None:
            eid = db.ejecutar(
                "INSERT INTO entidades (id_tipo_documento, id_pais, documento, tipo_entidad, fecha_registro, rubro) "
                "VALUES (%s, %s, %s, %s, %s, %s)", (ruc, pais, doc, "NATURAL" if natural else "JURIDICA", hoy, rubro))
            if natural:
                nombres, pat, mat = dividir_nombre(nombre)
                db.ejecutar("INSERT INTO personas_naturales (id_entidades, nombre, ape_pat, ape_mat) VALUES (%s, %s, %s, %s)",
                            (eid, nombres, pat, mat))
            else:
                db.ejecutar("INSERT INTO personas_juridicas (id_entidades, razon_social) VALUES (%s, %s)", (eid, nombre))
            db.contar("entidades nuevas")
        descripcion = f"{sancion} - Resolucion consentida en {anio}" + (" - Contencioso administrativo: Si" if str(contencioso).startswith("S") else "")
        if db.uno("SELECT id FROM historial_manchas WHERE id_entidades = %s AND id_tipo_lista = %s", (eid, lista)) is None:
            db.ejecutar("INSERT INTO historial_manchas (id_entidades, id_tipo_lista, descripcion, fecha_registro) VALUES (%s, %s, %s, %s)",
                        (eid, lista, descripcion, hoy))
            db.contar("sanciones registradas")
    print(f"   filas del Excel: {len(filas)}")


# ----------------------------------------------------------------------------- politicos

def leer_politicos(args):
    if args.politicos_json:
        return json.load(open(args.politicos_json, encoding="utf-8"))
    req = urllib.request.Request(args.ml_url.rstrip("/") + "/api/v1/personas?limite=500&con_casos=true",
                                 headers={"X-Internal-Key": args.ml_clave or ""})
    return json.load(urllib.request.urlopen(req, timeout=60))


def cargar_politicos(db: Base, args):
    personas = leer_politicos(args)
    print(f"\n== Politicos: {len(personas)} personas leidas del ML")
    pais = id_pais(db)
    sin_doc = id_tipo_doc(db, "SIN DOCUMENTO", "Sin documento registrado (figura publica vigilada)")
    lista_pep = id_lista_like(db, r"^Peru - Personas Politicamente Expuestas")
    if lista_pep is None:
        sys.exit("Falta la lista 'Perú - Personas Políticamente Expuestas': ejecuta primero el comando 'listas'.")
    lista_noticias = id_lista(db, LISTA_NOTICIAS)
    if lista_noticias is None:
        lista_noticias = db.ejecutar(
            "INSERT INTO tipo_lista (nombre, descripcion, fuente, alcance, periodicidad, es_pep, id_grupo) VALUES (%s,%s,%s,%s,%s,0,%s)",
            (LISTA_NOTICIAS, "Noticias de prensa nacional sobre personas vigiladas (acusaciones, denuncias, investigaciones, sentencias). "
             "Pendientes de revision humana salvo que se indique lo contrario.", "Diarios nacionales (vigilancia automatica)",
             "NACIONAL", "Diaria", id_grupo(db, GRUPOS[5][1])))
        db.contar("lista Noticias creada")
    hoy = date.today()
    for p in personas:
        doc = p.get("numero_documento") or f"SIN-DOC-{str(p['id']).replace('-', '')[:8].upper()}"
        tipo_doc = id_tipo_doc(db, p["tipo_documento"]) if p.get("numero_documento") and p.get("tipo_documento") else sin_doc
        eid = db.uno("SELECT id FROM entidades WHERE id_tipo_documento = %s AND documento = %s", (tipo_doc, doc))
        if eid is None:
            eid = db.ejecutar("INSERT INTO entidades (id_tipo_documento, id_pais, documento, tipo_entidad, fecha_registro, alias) "
                              "VALUES (%s, %s, %s, 'NATURAL', %s, %s)",
                              (tipo_doc, pais, doc, hoy, ", ".join(p.get("nombres_alternativos") or []) or None))
            db.ejecutar("INSERT INTO personas_naturales (id_entidades, nombre, ape_pat) VALUES (%s, %s, %s)",
                        (eid, p["nombres"], p["apellidos"]))
            db.contar("politicos nuevos")
        if p.get("es_pep") and db.uno("SELECT id FROM historial_manchas WHERE id_entidades=%s AND id_tipo_lista=%s", (eid, lista_pep)) is None:
            db.ejecutar("INSERT INTO historial_manchas (id_entidades, id_tipo_lista, descripcion, fecha_registro, cargo, tipo_pep) "
                        "VALUES (%s, %s, %s, %s, %s, 'NACIONAL')",
                        (eid, lista_pep, "Persona expuesta politicamente (figura publica vigilada en prensa)", hoy, p.get("cargo_pep")))
            db.contar("marcas PEP")
        for caso in p.get("casos") or []:
            estado = str(caso.get("estado_revision", "")).upper()
            if estado == "DESCARTADO" or not caso.get("url_fuente"):
                continue
            if db.uno("SELECT id FROM historial_manchas WHERE id_entidades=%s AND id_tipo_lista=%s AND link=%s",
                      (eid, lista_noticias, caso["url_fuente"])) is not None:
                continue
            etiqueta = "Pendiente de revision" if estado == "PENDIENTE" else estado.capitalize()
            fecha = (caso.get("creado_en") or str(hoy))[:10]
            db.ejecutar("INSERT INTO historial_manchas (id_entidades, id_tipo_lista, descripcion, link, fecha_registro, cargo) "
                        "VALUES (%s, %s, %s, %s, %s, %s)",
                        (eid, lista_noticias, f"[{etiqueta}] {str(caso.get('tipo','')).capitalize()}: {caso.get('resumen','')}",
                         caso["url_fuente"], fecha, caso.get("categoria_delito")))
            db.contar("noticias registradas")


# ----------------------------------------------------------------------------- principal

def _archivo(carpeta, nombre):
    buscado = unicodedata.normalize("NFC", nombre)
    for f in os.listdir(carpeta):
        if unicodedata.normalize("NFC", f) == buscado:
            return os.path.join(carpeta, f)
    sys.exit(f"No encuentro '{nombre}' en {carpeta}")


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("comandos", nargs="+", choices=["listas", "sancionadas", "politicos", "todo"])
    ap.add_argument("--carpeta", help="carpeta con los Excel (para listas y sancionadas)")
    ap.add_argument("--aplicar", action="store_true", help="escribe de verdad; sin esto solo simula")
    ap.add_argument("--env", default=os.path.join(os.path.dirname(__file__), "..", ".env"), help="archivo .env del backend")
    ap.add_argument("--db-url"), ap.add_argument("--usuario"), ap.add_argument("--clave")
    ap.add_argument("--politicos-json", help="JSON de personas con casos exportado del ML")
    ap.add_argument("--ml-url", help="URL del ML (usa GET /api/v1/personas?con_casos=true)")
    ap.add_argument("--ml-clave", default=os.environ.get("INTERNAL_API_KEY"), help="clave interna del ML")
    args = ap.parse_args()

    comandos = set(args.comandos)
    if "todo" in comandos:
        comandos |= {"listas", "sancionadas"} | ({"politicos"} if (args.politicos_json or args.ml_url) else set())
    if comandos & {"listas", "sancionadas"} and not args.carpeta:
        sys.exit("--carpeta es obligatoria para 'listas' y 'sancionadas'")
    if "politicos" in comandos and not (args.politicos_json or args.ml_url):
        sys.exit("'politicos' necesita --politicos-json o --ml-url")

    conexion = conectar(args)
    db = Base(conexion)
    try:
        if "listas" in comandos: cargar_listas(db, args.carpeta)
        if "sancionadas" in comandos: cargar_sancionadas(db, args.carpeta)
        if "politicos" in comandos: cargar_politicos(db, args)
        print("\n== Resumen:", db.stats or "sin cambios (ya estaba todo cargado)")
        if args.aplicar:
            conexion.commit(); print("APLICADO: cambios guardados.")
        else:
            conexion.rollback(); print("SIMULACION: no se guardo nada. Repite con --aplicar para escribir.")
    except Exception:
        conexion.rollback(); raise
    finally:
        conexion.close()


if __name__ == "__main__":
    main()
